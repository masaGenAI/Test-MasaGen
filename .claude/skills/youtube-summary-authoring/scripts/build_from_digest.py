"""チャンネル・ダイジェスト（source-digest-report 形式の xlsx）から、Book-Summary の YouTube 枠に入れる
骨組み（メタ情報・注記・棚・出典ツリー・チャンネル総括）と、要約執筆用の入力を作る。

usage: python3 -I build_from_digest.py <src_dir> <out_dir> <channels.json> [--hub <booksummaryhub.html>]
  src_dir/<フォルダ>/<名前>_ダイジェスト.xlsx を読む。
  out_dir に skel.json（id→骨組み）, inputs/<id>.json（執筆用）, tree.json, digest.json を書く。
"""
import json, os, re, sys, glob
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
# --hub <booksummaryhub.html> を渡すと、すでに収録済みの動画（URL の動画IDで判定）を飛ばす
HAVE = set()
if "--hub" in sys.argv:
    _h = open(sys.argv[sys.argv.index("--hub") + 1], encoding="utf-8").read()
    HAVE = set(re.findall(r'"url":"https://(?:www\.)?youtube\.com/(?:watch\?v=|shorts/)([\w-]{11})', _h))
skipped = 0
os.makedirs(OUT + "/inputs", exist_ok=True)

# チャンネル定義は JSON（フォルダ名 → id・表示名・副題・下位ノード列と対応・テーマ→棚）
CH = json.load(open(sys.argv[3], encoding="utf-8"))

def g(d, *keys):
    for k in keys:
        for kk, v in d.items():
            if kk == k or kk.startswith(k):
                if v not in (None, ""):
                    return str(v).strip()
    return ""

def minutes(d):
    v = g(d, "尺（分）", "長さ（分）", "動画尺（分）")
    if v:
        try: return f"{round(float(v))}分"
        except ValueError: pass
    v = g(d, "尺（秒）")
    if v:
        try:
            s = int(float(v)); return f"{s}秒" if s < 120 else f"{round(s/60)}分"
        except ValueError: pass
    return ""

def vid(u):
    m = re.search(r"(?:v=|shorts/|youtu\.be/)([\w-]{11})", u or ""); return m.group(1) if m else None

skel, tree, digest, themes = {}, [], {}, {}
for f in sorted(glob.glob(SRC + "/*/*.xlsx")):
    folder = f.split("/")[-2]
    if folder not in CH: continue
    c = CH[folder]
    wb = openpyxl.load_workbook(f, read_only=True)
    ws = [w for w in wb.worksheets if w.title.startswith("01")][0]
    rows = list(ws.iter_rows(values_only=True)); H = [str(h) for h in rows[0]]
    kids = {}
    for r in rows[1:]:
        if not r or not r[0]: continue
        d = {H[i]: (r[i] if i < len(r) else None) for i in range(len(H))}
        no = int(str(d["No"]).strip()); url = g(d, "URL"); date = str(g(d, "公開日"))[:10]
        if vid(url) in HAVE:
            skipped += 1; continue
        theme = g(d, "テーマ", "所属再生リスト：テーマ")
        sub = None
        if c.get("child_col"):
            key = g(d, c["child_col"])
            sub = c["children"][key][0]; kids[sub] = c["children"][key]
        rid = f"yt_{c['id']}_{sub or 'main'}_{no}"
        title = re.sub(r"[|｜]\s*$", "", g(d, "動画タイトル", "記事タイトル")).strip()
        notes = []
        ev = g(d, "エビデンス")
        if ev: notes.append({"h": "エビデンス（数値・出典）", "b": ev})
        book = g(d, "原典", "解説対象の書籍")
        if book and book not in ("なし", "-", "—"): notes.append({"h": "解説対象の書籍", "b": book})
        src = g(d, "扱う思考法・概念")
        if src: notes.append({"h": "依拠する研究・思想・情報源", "b": src})
        q = g(d, "名言")
        if q: notes.append({"h": "取り上げた名言", "b": q + (f"（{g(d, '発言者')}" + (f"／{g(d, '発言者の属性')}" if g(d, '発言者の属性') else "") + "）" if g(d, "発言者") else "")})
        sp = g(d, "話者・登壇者")
        if sp: notes.append({"h": "話者・登壇者", "b": sp})
        org = g(d, "出典（元番組・元動画）")
        if org: notes.append({"h": "元の番組・動画", "b": org})
        caveat = "強度 " + (g(d, "エビデンス強度", "強度") or "—")
        why = g(d, "強度判定の根拠")
        if why: caveat += "／" + why
        pr = g(d, "PR・宣伝", "動画内告知・PR")
        if pr and not pr.startswith("なし"): caveat += "／動画内の宣伝・告知: " + pr.replace("あり：", "")
        tie = g(d, "企業・自治体の提供案件")
        if tie and tie != "なし": caveat += "／提供案件: " + tie
        if c.get("basis") == "概要欄のみ": caveat += "／本編は焼き込み字幕のため、内容は概要欄（チャンネル運営者による日本語要約）に依拠"
        notes.append({"h": "エビデンス強度と留保", "b": caveat})
        mn = minutes(d)
        rec = {"id": rid, "title": title, "author": "", "publisher": "", "year": date[:4],
               "url": url, "srcNote": " ・ ".join(x for x in [mn, date] if x),
               "srcPaths": [["yt", c["id"]] + ([sub] if sub else [])], "themes": [theme],
               "shelf": c["shelf"][theme], "notes": notes}
        # 書籍解説チャンネル: 紹介書籍と著者
        if folder == "アバタロー" and book:
            m = re.match(r"(.+?)／(.+?)(?:（.*）)?$", book)
            if m: rec["refBooks"], rec["author"] = m.group(1).strip(), m.group(2).strip()
        if folder == "学識サロン":
            fmt = g(d, "形式")
            if fmt == "書籍1冊の解説" and book:
                m = re.match(r"『(.+?)』\s*(.*)$", book)
                if m: rec["refBooks"], rec["author"] = m.group(1), m.group(2).strip()
            elif fmt in ("総集編（複数書籍）",) and book:
                rec["refBooks"] = "／".join(re.findall(r"『(.+?)』", book)) or book
            else:
                rec["kind"] = "theme"
        skel[rid] = rec
        inp = {"id": rid, "channel": c["title"], "title": title, "date": date, "length": mn,
               "format": g(d, "形式") or ("ショート" if "/shorts/" in url else "長尺"),
               "theme": theme, "claim": g(d, "論点"), "summary": g(d, "動画概要", "動画の概要", "サマリ"),
               "chapters": g(d, "章立て"), "evidence": ev, "implication": g(d, "実務示唆"),
               "book": book, "quote": q, "speaker": g(d, "発言者") or sp, "speaker_attr": g(d, "発言者の属性"),
               "concepts": src or g(d, "概念タグ"), "strength": g(d, "エビデンス強度", "強度"), "pr": pr,
               "source_basis": c.get("basis", "字幕全文＋概要欄")}
        fmt_short = inp["format"] == "ショート" or "/shorts/" in url
        inp["tier"] = "short" if fmt_short else ("full" if len(inp["summary"]) >= 150 or len(inp["chapters"]) >= 60 else "desc")
        json.dump(inp, open(f"{OUT}/inputs/{rid}.json", "w"), ensure_ascii=False, indent=1)
    node = {"id": c["id"], "title": c["title"], "sub": c["sub"], "icon": "folder"}
    if kids:
        order = list(c["children"].values())
        node["children"] = [{"id": k[0], "title": k[1], "sub": k[2]} for k in order if k[0] in kids]
    tree.append(node)
    # チャンネル総括（00_サマリ）
    srows = [[x for x in r if x is not None] for r in wb["00_サマリ"].iter_rows(values_only=True)]
    desc = next((str(r[0]).split("チャンネル概要:")[1].strip() for r in srows if r and str(r[0]).startswith("チャンネル概要:")), "")
    head = str(srows[0][0]); scope = str(srows[1][0]) if len(srows) > 1 else ""
    pts, note = [], ""
    for r in srows:
        if r and str(r[0]).isdigit() and int(str(r[0])) <= 20 and len(r) >= 4 and not str(r[1]).startswith("="):
            b = f"{r[2]}。→ {r[3]}" + (f"（関連No. {r[4]}）" if len(r) > 4 else "")
            pts.append({"h": str(r[1]), "b": b})
        if r and str(r[0]).startswith("※注記"): note = str(r[0]).replace("※注記: ", "")
    lim = []
    ws4 = [w for w in wb.worksheets if w.title.startswith("04")][0]
    on = False
    for r in ws4.iter_rows(values_only=True):
        t = [str(x) for x in r if x]
        if not t: continue
        if "制約" in t[0] and "留意" in t[0]:
            on = True
            if len(t) > 1: lim.append(" ".join(t[1:]))
            continue
        if on:
            if re.match(r"^\d+\.", t[0]): break
            lim.append(" ".join(t))
    # 制約は (a)(b)… の項目だけを残す（除外動画の一覧やリンク疎通の記録は載せない）
    lim = [y.strip() for x in lim for y in re.split(r"(?:(?<=。)|\s)\s*(?=\([b-z]\)\s?)", x)]
    lim = [re.sub(r"^(\([a-z]\)\s*)(apple-to-apple:\s*)?", r"\1", x) for x in lim
           if re.match(r"^\([a-z]\)", x) and "リンク疎通" not in x]
    # テーマの読み筋（02_論点集約: テーマ・本数・構成比・代表・結論・代表エビデンス・実務示唆）
    tn = {}
    for r in wb["02_論点集約"].iter_rows(values_only=True):
        used = {x["themes"][0] for x in skel.values() if x["srcPaths"][0][1] == c["id"]}
        if r and r[0] in used and len(r) >= 7 and str(r[1] or "").startswith("=COUNTIF") and r[4] and str(r[0]) not in tn:
            tn[str(r[0])] = {"c": str(r[4]), "e": str(r[5] or ""), "a": str(r[6] or "")}
    for p in ["yt/" + c["id"]] + ["yt/" + c["id"] + "/" + k["id"] for k in node.get("children", [])]:
        themes[p] = tn
    digest["yt/" + c["id"]] = {"summary": (desc + " " if desc else "") + head, "scope": scope, "points": pts,
                               "caveats": [x for x in [note] + lim if x]}

json.dump(skel, open(OUT + "/skel.json", "w"), ensure_ascii=False, indent=1)
json.dump(themes, open(OUT + "/theme_note.json", "w"), ensure_ascii=False, indent=1)
json.dump(tree, open(OUT + "/tree.json", "w"), ensure_ascii=False, indent=1)
json.dump(digest, open(OUT + "/digest.json", "w"), ensure_ascii=False, indent=1)
print(len(skel), "records;", len(tree), "channels;", skipped, "already in hub (skipped)")
